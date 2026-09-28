from __future__ import annotations

from collections import Counter, defaultdict
import hashlib
from pathlib import Path
import re
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook


HEADER_ALIASES = {
    "billingperiod": "period",
    "customerno": "customer",
    "customername": "customer_name",
    "partno": "part",
    "serialno": "serial",
    "pl2": "pl2",
    "pl4": "pl4",
    "pl6": "pl6",
    "contract": "contract",
    "bwdv": "bw_dv",
    "bwa4dv": "bw_a4_dv",
    "bwa3dv": "bw_a3_dv",
    "fullcolourdv": "full_colour_dv",
    "coloura4dv": "colour_a4_dv",
    "coloura3dv": "colour_a3_dv",
    "totaldv": "total_dv",
    "bwcv": "bw_cv",
    "bwa4cv": "bw_a4_cv",
    "bwa3cv": "bw_a3_cv",
    "fullcolourcv": "full_colour_cv",
    "coloura4cv": "colour_a4_cv",
    "coloura3cv": "colour_a3_cv",
    "totalcv": "total_cv",
    "bwrev": "bw_rev",
    "bwa4rev": "bw_a4_rev",
    "bwa3rev": "bw_a3_rev",
    "fullcolourrev": "full_colour_rev",
    "coloura4rev": "colour_a4_rev",
    "coloura3rev": "colour_a3_rev",
    "totalrev": "total_rev",
    "currency": "currency",
    "billto": "bill_to",
    "billtolocationname": "bill_to_location",
    "shipto": "ship_to",
    "shiptolocationname": "ship_to_location",
    "costcenter": "cost_center",
}


def _normalized(value: Any) -> str:
    if value is None:
        return "<NULL>"
    return str(value).strip()


def _fingerprint(row: Mapping[str, Any], fields: Sequence[str] | None = None) -> tuple[str, ...]:
    selected = fields if fields is not None else sorted(row)
    return tuple(_normalized(row.get(field)) for field in selected)


def _row_digest(row: Mapping[str, Any]) -> bytes:
    payload = "\x1f".join(
        f"{field}\x1e{_normalized(row.get(field))}" for field in sorted(row)
    ).encode("utf-8", errors="replace")
    return hashlib.blake2b(payload, digest_size=16).digest()


def _canonical_headers(values: Sequence[Any]) -> list[str]:
    headers: list[str] = []
    seen: Counter[str] = Counter()
    for index, value in enumerate(values, start=1):
        raw = "" if value is None else str(value).strip()
        if raw == "&":
            base = "equipment_key"
        else:
            compact = re.sub(r"[^a-z0-9]", "", raw.lower())
            base = HEADER_ALIASES.get(compact, f"source_{index}")
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def iter_workbook_rows(
    workbook_path: Path | str,
    sheet_name: str,
    *,
    header_row: int,
) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        header_values = next(
            sheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            )
        )
        last_nonempty = max(
            (index for index, value in enumerate(header_values) if value is not None),
            default=-1,
        )
        headers = _canonical_headers(header_values[: last_nonempty + 1])
        for values in sheet.iter_rows(
            min_row=header_row + 1,
            max_col=len(headers),
            values_only=True,
        ):
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            yield dict(zip(headers, values))
    finally:
        workbook.close()


def profile_rows(
    rows: Iterable[Mapping[str, Any]],
    *,
    key_fields: Sequence[str],
    measure_fields: Sequence[str] = (),
) -> dict[str, Any]:
    exact_counts: Counter[bytes] = Counter()
    exact_periods: dict[bytes, str] = {}
    key_counts: Counter[tuple[str, ...]] = Counter()
    key_fingerprints: dict[tuple[str, ...], set[bytes]] = defaultdict(set)
    key_measures: dict[tuple[str, ...], set[tuple[str, ...]]] = defaultdict(set)
    key_periods: dict[tuple[str, ...], str] = {}
    row_count = 0

    for row in rows:
        row_count += 1
        exact = _row_digest(row)
        key = _fingerprint(row, key_fields)
        period = _normalized(row.get("period"))
        exact_counts[exact] += 1
        exact_periods[exact] = period
        key_counts[key] += 1
        key_fingerprints[key].add(exact)
        key_measures[key].add(_fingerprint(row, measure_fields))
        key_periods[key] = period

    exact_duplicates = {fp: count for fp, count in exact_counts.items() if count > 1}
    duplicate_by_period: Counter[str] = Counter()
    for fingerprint, count in exact_duplicates.items():
        duplicate_by_period[exact_periods[fingerprint]] += count - 1

    duplicate_keys = {key: count for key, count in key_counts.items() if count > 1}
    key_duplicate_by_period: Counter[str] = Counter()
    for key, count in duplicate_keys.items():
        key_duplicate_by_period[key_periods[key]] += count - 1
    return {
        "row_count": row_count,
        "exact_duplicate_excess_rows": sum(count - 1 for count in exact_duplicates.values()),
        "exact_duplicate_affected_rows": sum(exact_duplicates.values()),
        "exact_duplicate_groups": len(exact_duplicates),
        "exact_duplicate_excess_by_period": dict(sorted(duplicate_by_period.items())),
        "max_exact_copies": max(exact_duplicates.values(), default=1),
        "key_duplicate_excess_rows": sum(count - 1 for count in duplicate_keys.values()),
        "key_duplicate_affected_rows": sum(duplicate_keys.values()),
        "key_duplicate_groups": len(duplicate_keys),
        "conflicting_key_groups": sum(
            len(key_fingerprints[key]) > 1 for key in duplicate_keys
        ),
        "same_measure_key_groups": sum(
            len(key_measures[key]) == 1 for key in duplicate_keys
        ),
        "conflicting_measure_key_groups": sum(
            len(key_measures[key]) > 1 for key in duplicate_keys
        ),
        "key_duplicate_excess_by_period": dict(sorted(key_duplicate_by_period.items())),
        "max_key_copies": max(duplicate_keys.values(), default=1),
    }


def compare_period_rows(
    left_rows: Iterable[Mapping[str, Any]],
    right_rows: Iterable[Mapping[str, Any]],
    *,
    period: Any,
    fields: Sequence[str],
) -> dict[str, int]:
    target = _normalized(period)
    left = Counter(
        _fingerprint(row, fields)
        for row in left_rows
        if _normalized(row.get("period")) == target
    )
    right = Counter(
        _fingerprint(row, fields)
        for row in right_rows
        if _normalized(row.get("period")) == target
    )
    matched = left & right
    matched_rows = sum(matched.values())
    left_rows_count = sum(left.values())
    right_rows_count = sum(right.values())
    return {
        "left_rows": left_rows_count,
        "right_rows": right_rows_count,
        "matched_rows": matched_rows,
        "left_only_rows": left_rows_count - matched_rows,
        "right_only_rows": right_rows_count - matched_rows,
    }
