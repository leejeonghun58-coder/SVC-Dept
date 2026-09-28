from __future__ import annotations

import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


path = Path(sys.argv[1])
workbook = load_workbook(path, read_only=True, data_only=False)
sheet = workbook[workbook.sheetnames[0]]
headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
index = {str(value): position for position, value in enumerate(headers)}

months: Counter[str] = Counter()
sr_min = None
sr_max = None
service_date_sentinel = 0
qty_min = None
qty_max = None

for row in sheet.iter_rows(min_row=2, values_only=True):
    month = row[index["월도"]]
    if month is not None:
        months[str(month).strip()] += 1
    sr_date = row[index["SR Date"]]
    if isinstance(sr_date, datetime):
        sr_min = sr_date if sr_min is None else min(sr_min, sr_date)
        sr_max = sr_date if sr_max is None else max(sr_max, sr_date)
    service_date = row[index["Service Date"]]
    if isinstance(service_date, (int, float)) and service_date > 100000:
        service_date_sentinel += 1
    qty = row[index["Qty"]]
    if isinstance(qty, (int, float)):
        qty_min = qty if qty_min is None else min(qty_min, qty)
        qty_max = qty if qty_max is None else max(qty_max, qty)

workbook.close()
print(
    json.dumps(
        {
            "months": dict(sorted(months.items())),
            "sr_date_min": sr_min.isoformat() if sr_min else None,
            "sr_date_max": sr_max.isoformat() if sr_max else None,
            "service_date_values_over_100000": service_date_sentinel,
            "qty_min": qty_min,
            "qty_max": qty_max,
        },
        ensure_ascii=True,
    )
)
