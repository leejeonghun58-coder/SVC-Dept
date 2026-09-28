from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADER_SCAN_ROWS = 30
UNIQUE_CAP = 10_000
SAMPLE_CAP = 5


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def mask_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value)
    text = safe_text(value)
    if not text:
        return ""
    if len(text) <= 2:
        return "*" * len(text)
    return f"{text[0]}{'*' * min(len(text) - 2, 6)}{text[-1]}"


def type_name(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    return "text"


def normalized_cell(value: Any) -> str:
    if value is None:
        return "<NULL>"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return "<NAN>"
    return str(value).strip()


def numeric_like(text: str) -> bool:
    if not text:
        return False
    cleaned = text.replace(",", "").replace(" ", "")
    return bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?", cleaned))


def detect_header(rows: list[tuple[Any, ...]]) -> tuple[int, list[str]]:
    best_index = 0
    best_score = -1.0
    for index, row in enumerate(rows):
        values = [safe_text(value) for value in row]
        nonempty = [value for value in values if value]
        if not nonempty:
            continue
        text_count = sum(not numeric_like(value) for value in nonempty)
        unique_count = len(set(nonempty))
        score = len(nonempty) + 0.5 * text_count + 0.25 * unique_count
        if score > best_score:
            best_score = score
            best_index = index
    header_values = [safe_text(value) for value in rows[best_index]]
    last_nonempty = max((i for i, value in enumerate(header_values) if value), default=-1)
    header_values = header_values[: last_nonempty + 1]
    seen: Counter[str] = Counter()
    headers: list[str] = []
    for column_index, value in enumerate(header_values, start=1):
        base = value or f"unnamed_{column_index}"
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return best_index + 1, headers


def profile_sheet(workbook_path: Path, sheet_name: str, sheet_state: str) -> dict[str, Any]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    ws = wb[sheet_name]
    scan_rows: list[tuple[Any, ...]] = []
    for row in ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True):
        scan_rows.append(tuple(row))
    header_row, headers = detect_header(scan_rows)
    column_count = len(headers)
    stats: list[dict[str, Any]] = [
        {
            "column": header,
            "nonempty": 0,
            "blank": 0,
            "types": Counter(),
            "unique_values": set(),
            "unique_capped": False,
            "samples": [],
            "numeric_text": 0,
            "zero_count": 0,
            "negative_count": 0,
        }
        for header in headers
    ]
    exact_hashes: set[bytes] = set()
    exact_duplicate_rows = 0
    data_rows = 0
    completely_blank_rows = 0
    formula_cells = 0

    for cells in ws.iter_rows(min_row=header_row + 1, max_col=column_count):
        values = [cell.value for cell in cells]
        if all(value is None or safe_text(value) == "" for value in values):
            completely_blank_rows += 1
            continue
        data_rows += 1
        row_key = "\x1f".join(normalized_cell(value) for value in values).encode("utf-8", errors="replace")
        row_hash = hashlib.blake2b(row_key, digest_size=16).digest()
        if row_hash in exact_hashes:
            exact_duplicate_rows += 1
        else:
            exact_hashes.add(row_hash)

        for index, cell in enumerate(cells):
            value = cell.value
            stat = stats[index]
            if cell.data_type == "f":
                formula_cells += 1
            if value is None or safe_text(value) == "":
                stat["blank"] += 1
                continue
            stat["nonempty"] += 1
            kind = type_name(value)
            stat["types"][kind] += 1
            normalized = normalized_cell(value)
            if not stat["unique_capped"]:
                stat["unique_values"].add(normalized)
                if len(stat["unique_values"]) > UNIQUE_CAP:
                    stat["unique_capped"] = True
                    stat["unique_values"].clear()
            if len(stat["samples"]) < SAMPLE_CAP:
                masked = mask_value(value)
                if masked and masked not in stat["samples"]:
                    stat["samples"].append(masked)
            if isinstance(value, str) and numeric_like(value):
                stat["numeric_text"] += 1
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                if value == 0:
                    stat["zero_count"] += 1
                elif value < 0:
                    stat["negative_count"] += 1

    columns: list[dict[str, Any]] = []
    for stat in stats:
        total = data_rows
        blank = int(stat["blank"])
        columns.append(
            {
                "column": stat["column"],
                "nonempty": int(stat["nonempty"]),
                "blank": blank,
                "blank_rate": round(blank / total, 6) if total else None,
                "types": dict(stat["types"]),
                "unique_count": f">{UNIQUE_CAP}" if stat["unique_capped"] else len(stat["unique_values"]),
                "samples_masked": stat["samples"],
                "numeric_text": int(stat["numeric_text"]),
                "zero_count": int(stat["zero_count"]),
                "negative_count": int(stat["negative_count"]),
            }
        )

    result = {
        "name": sheet_name,
        "state": sheet_state,
        "reported_dimension": ws.calculate_dimension(),
        "header_row": header_row,
        "headers": headers,
        "data_rows_nonblank": data_rows,
        "blank_rows_after_header": completely_blank_rows,
        "column_count": column_count,
        "formula_cells": formula_cells,
        "exact_duplicate_rows": exact_duplicate_rows,
        "exact_duplicate_rate": round(exact_duplicate_rows / data_rows, 6) if data_rows else None,
        "columns": columns,
    }
    wb.close()
    return result


def profile_workbook(path: Path) -> dict[str, Any]:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    stat = path.stat()
    wb = load_workbook(path, read_only=True, data_only=False)
    sheet_meta = [(ws.title, ws.sheet_state) for ws in wb.worksheets]
    wb.close()
    sheets = [profile_sheet(path, name, state) for name, state in sheet_meta]
    return {
        "file_name": path.name,
        "source_path": str(path),
        "size_bytes": stat.st_size,
        "modified_local": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        "sha256": digest.hexdigest().upper(),
        "sheets": sheets,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    profiles = [profile_workbook(path.resolve()) for path in args.workbooks]
    payload = {"generated_at": datetime.now().isoformat(timespec="seconds"), "workbooks": profiles}
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
